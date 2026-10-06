from __future__ import annotations

import re
from dataclasses import dataclass, replace
from html.parser import HTMLParser
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.token import Token

from .html_builder import (
    DEFAULT_MERMAID_CDN,
    EMOJI_PATTERN,
    convert_mermaid_fences_to_html,
    extract_title,
    read_text_safely,
    strip_emojis,
    strip_front_matter,
)

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from docx.shared import Length
    from docx.text.paragraph import Paragraph


DEFAULT_DOCX_FONT = "Vazirmatn"
DEFAULT_DOCX_FONT_FALLBACK = "Tahoma"
DEFAULT_DOCX_CODE_FONT = "Consolas"
DEFAULT_DOCX_EMOJI_FONT = "Segoe UI Emoji"

HEADING_SIZES = {1: 22, 2: 18, 3: 15, 4: 13, 5: 12, 6: 11}

PAGE_FORMAT_DIMENSIONS: dict[str, tuple[float, float]] = {
    "A4": (210.0, 297.0),
    "LETTER": (215.9, 279.4),
    "LEGAL": (215.9, 355.6),
}


UNIT_CONVERTERS: dict[str, Any] = {
    "mm": Mm,
    "cm": Cm,
    "in": Inches,
    "inch": Inches,
    "pt": Pt,
    "px": lambda v: Inches(v / 96.0),
}


def parse_docx_length(margin_str: str) -> Length:
    """Parse length string (e.g. '5mm', '15mm', '0.5in', '1.5cm', '20px') into docx Length."""
    m = re.match(r"^([\d.]+)\s*(mm|cm|in|inch|pt|px)?$", (margin_str or "").strip().lower())
    if not m:
        return Mm(15.0)
    try:
        val = float(m.group(1))
    except ValueError:
        return Mm(15.0)
    converter = UNIT_CONVERTERS.get(m.group(2) or "mm", Mm)
    return converter(val)



@dataclass(frozen=True)
class DocxBuildOptions:
    source_path: Path
    font_family: str = DEFAULT_DOCX_FONT
    font_size_pt: int = 12
    rtl: bool = True
    mermaid_images: dict[str, Path] | None = None  # mermaid_code -> png path
    image_scale: int = 3                  # device scale factor used for the PNGs
    image_min_width_inches: float = 4.0   # minimum image width in document
    image_max_width_inches: float = 6.5   # maximum image width in document
    include_page_numbers: bool = True
    highlight_code: bool = True
    page_format: str = "A4"
    margin: str = "15mm"
    landscape: bool = False
    strip_emojis: bool = False


@dataclass(frozen=True)
class DocxBuildResult:
    title: str
    has_mermaid: bool
    mermaid_blocks: list[str]  # raw mermaid code blocks (for image rendering)


def _set_paragraph_rtl(paragraph: "Paragraph") -> None:
    """Apply right-to-left direction to a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is not None:
        pPr.remove(bidi)
    bidi = pPr.makeelement(qn("w:bidi"), {})
    pPr.append(bidi)


def _set_paragraph_ltr(paragraph: "Paragraph") -> None:
    """Explicitly apply left-to-right direction to a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is not None:
        pPr.remove(bidi)
    bidi = pPr.makeelement(qn("w:bidi"), {qn("w:val"): "0"})
    pPr.append(bidi)


def _set_run_rtl(run) -> None:
    """Apply right-to-left text formatting to a run."""
    rPr = run._r.get_or_add_rPr()
    rtl = rPr.find(qn("w:rtl"))
    if rtl is not None:
        rPr.remove(rtl)
    rtl = rPr.makeelement(qn("w:rtl"), {})
    rPr.append(rtl)


def _set_run_ltr(run) -> None:
    """Explicitly apply left-to-right text formatting to a run."""
    rPr = run._r.get_or_add_rPr()
    rtl = rPr.find(qn("w:rtl"))
    if rtl is not None:
        rPr.remove(rtl)
    rtl = rPr.makeelement(qn("w:rtl"), {qn("w:val"): "0"})
    rPr.append(rtl)


def _set_run_size_cs(run, font_size_pt: int) -> None:
    """Set complex-script font size (szCs) used for Arabic/Persian text in Word."""
    rPr = run._r.get_or_add_rPr()
    # Word stores size in half-points
    half_pts = str(int(font_size_pt * 2))
    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = rPr.makeelement(qn("w:szCs"), {})
        rPr.append(szCs)
    szCs.set(qn("w:val"), half_pts)


def _set_run_bold_cs(run) -> None:
    """Set complex-script bold (bCs) used for Arabic/Persian text in Word."""
    rPr = run._r.get_or_add_rPr()
    bCs = rPr.find(qn("w:bCs"))
    if bCs is None:
        bCs = rPr.makeelement(qn("w:bCs"), {})
        rPr.append(bCs)


def _set_run_italic_cs(run) -> None:
    """Set complex-script italic (iCs) used for Arabic/Persian text in Word."""
    rPr = run._r.get_or_add_rPr()
    iCs = rPr.find(qn("w:iCs"))
    if iCs is None:
        iCs = rPr.makeelement(qn("w:iCs"), {})
        rPr.append(iCs)


def _set_run_font(run, font_name: str, complex_script_font: str | None = None) -> None:
    """Set the font for both ASCII and complex script (Arabic/Persian)."""
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), complex_script_font or font_name)


def _set_document_default_rtl(document: "DocxDocument", font_name: str, font_size_pt: int) -> None:
    """Configure document-wide defaults for RTL Persian."""
    style = document.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(font_size_pt)

    # Set RTL on the Normal style paragraph properties
    style_element = style.element
    pPr = style_element.find(qn("w:pPr"))
    if pPr is None:
        pPr = style_element.makeelement(qn("w:pPr"), {})
        style_element.append(pPr)
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = pPr.makeelement(qn("w:bidi"), {})
        pPr.append(bidi)
    # Remove any jc - let bidi handle alignment naturally
    jc = pPr.find(qn("w:jc"))
    if jc is not None:
        pPr.remove(jc)

    # Set RTL on the Normal style run properties
    rPr = style_element.find(qn("w:rPr"))
    if rPr is None:
        rPr = style_element.makeelement(qn("w:rPr"), {})
        style_element.append(rPr)
    rtl = rPr.find(qn("w:rtl"))
    if rtl is None:
        rtl = rPr.makeelement(qn("w:rtl"), {})
        rPr.append(rtl)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), font_name)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)

    # Also set heading styles to RTL
    for heading_num in range(1, 7):
        style_name = f"Heading {heading_num}"
        try:
            h_style = document.styles[style_name]
            h_element = h_style.element
            h_pPr = h_element.find(qn("w:pPr"))
            if h_pPr is None:
                h_pPr = h_element.makeelement(qn("w:pPr"), {})
                h_element.append(h_pPr)
            h_bidi = h_pPr.find(qn("w:bidi"))
            if h_bidi is None:
                h_bidi = h_pPr.makeelement(qn("w:bidi"), {})
                h_pPr.append(h_bidi)
            # Remove jc from headings too
            h_jc = h_pPr.find(qn("w:jc"))
            if h_jc is not None:
                h_pPr.remove(h_jc)
        except KeyError:
            pass


@dataclass(frozen=True)
class RunStyle:
    font_name: str = DEFAULT_DOCX_FONT
    font_size_pt: int | None = None
    bold: bool = False
    italic: bool = False
    color: RGBColor | None = None
    code: bool = False
    rtl: bool = True


def _split_emojis(text: str) -> list[tuple[str, bool]]:
    """Split text into (substring, is_emoji) pairs."""
    result: list[tuple[str, bool]] = []
    last_end = 0
    for match in EMOJI_PATTERN.finditer(text):
        start, end = match.span()
        if start > last_end:
            result.append((text[last_end:start], False))
        result.append((match.group(0), True))
        last_end = end
    if last_end < len(text):
        result.append((text[last_end:], False))
    return result or [(text, False)]


def _apply_run_formatting(run, style: RunStyle) -> None:
    """Apply size, weight, slant, color, and bidi formatting to a run."""
    if style.font_size_pt is not None:
        run.font.size = Pt(style.font_size_pt)
        _set_run_size_cs(run, style.font_size_pt)
    if style.bold:
        run.bold = True
        _set_run_bold_cs(run)
    if style.italic:
        run.italic = True
        _set_run_italic_cs(run)
    if style.color is not None:
        run.font.color.rgb = style.color
    if style.rtl and not style.code:
        _set_run_rtl(run)
    else:
        _set_run_ltr(run)


def _add_single_run(paragraph: "Paragraph", text: str, style: RunStyle):
    """Add a single formatted run to a paragraph."""
    run = paragraph.add_run(text)
    f_name = style.font_name
    cs_font = None if style.code else f_name
    _set_run_font(run, f_name, complex_script_font=cs_font)
    _apply_run_formatting(run, style)
    return run


def _add_emoji_split_runs(
    paragraph: "Paragraph",
    text: str,
    base_style: RunStyle,
    base_font: str,
):
    """Split text by emojis and add individual runs with matching fonts."""
    last_run = None
    for piece, is_emoji in _split_emojis(text):
        style = replace(
            base_style,
            font_name=DEFAULT_DOCX_EMOJI_FONT if is_emoji else base_font,
            rtl=False if is_emoji else base_style.rtl,
        )
        last_run = _add_single_run(paragraph, piece, style)
    return last_run


def _add_styled_run(
    paragraph: "Paragraph",
    text: str,
    *,
    font_name: str = DEFAULT_DOCX_FONT,
    font_size_pt: int | None = None,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
    rtl: bool = True,
    color: RGBColor | None = None,
):
    """Add a run to the paragraph with styling and emoji font fallback."""
    if not text:
        return None
    style = RunStyle(
        font_name=DEFAULT_DOCX_CODE_FONT if code else font_name,
        font_size_pt=font_size_pt,
        bold=bold,
        italic=italic,
        color=color,
        code=code,
        rtl=rtl,
    )
    if not code and EMOJI_PATTERN.search(text):
        return _add_emoji_split_runs(paragraph, text, style, font_name)
    return _add_single_run(paragraph, text, style)


def _create_hyperlink_piece(hyperlink, paragraph: "Paragraph", text: str, style: RunStyle) -> None:
    """Create a styled run inside a Word hyperlink element."""
    from docx.oxml import parse_xml
    from docx.text.run import Run

    run_xml = parse_xml(r'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    hyperlink.append(run_xml)
    run = Run(run_xml, paragraph)
    run.text = text
    cs_font = style.font_name if not style.rtl else None
    _set_run_font(run, style.font_name, complex_script_font=cs_font)
    _apply_run_formatting(run, style)
    run.underline = True


def _add_hyperlink_run(
    paragraph: "Paragraph",
    text: str,
    href: str,
    *,
    font_name: str = DEFAULT_DOCX_FONT,
    font_size_pt: int | None = None,
    bold: bool = False,
    italic: bool = False,
    rtl: bool = True,
) -> None:
    """Add a native Word hyperlink to a paragraph with emoji support."""
    try:
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import parse_xml

        part = paragraph.part
        r_id = part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = parse_xml(
            f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}"/>'
        )
        paragraph._p.append(hyperlink)

        base_style = RunStyle(
            font_name=font_name,
            font_size_pt=font_size_pt,
            bold=bold,
            italic=italic,
            color=RGBColor(0x07, 0x58, 0x85),
            rtl=rtl,
        )
        for piece, is_emoji in _split_emojis(text):
            style = replace(
                base_style,
                font_name=DEFAULT_DOCX_EMOJI_FONT if is_emoji else font_name,
                rtl=False if is_emoji else rtl,
            )
            _create_hyperlink_piece(hyperlink, paragraph, piece, style)
    except Exception:
        _add_styled_run(
            paragraph,
            text,
            font_name=font_name,
            font_size_pt=font_size_pt,
            bold=bold,
            italic=italic,
            color=RGBColor(0x07, 0x58, 0x85),
            rtl=rtl,
        )


HTML_IMG_RE = re.compile(r"""<img\b(?P<attrs>[^>]*)>""", re.IGNORECASE)
ATTR_SRC_RE = re.compile(r"""src\s*=\s*['"]([^'"]+)['"]""", re.IGNORECASE)
ATTR_ALT_RE = re.compile(r"""alt\s*=\s*['"]([^'"]*)['"]""", re.IGNORECASE)


def _calculate_image_width(img_px_width: int, options: DocxBuildOptions) -> float:
    """Calculate target width in inches constrained by page format and margins."""
    css_width = img_px_width / 96.0
    fmt_key = (options.page_format or "A4").strip().upper()
    w_mm, h_mm = PAGE_FORMAT_DIMENSIONS.get(fmt_key, (210.0, 297.0))
    page_w_in = (h_mm if options.landscape else w_mm) / 25.4
    margin_in = parse_docx_length(options.margin).inches
    available_w = max(1.0, page_w_in - (2 * margin_in))
    max_w = min(options.image_max_width_inches, available_w) if options.image_max_width_inches != 6.5 else available_w
    target_w = min(css_width, max_w)
    return min(css_width, available_w) if target_w < 1.0 else target_w


def _resolve_image_file(src: str, source_path: Path) -> Path | None:
    """Resolve relative or absolute local image file path."""
    if not src or src.startswith(("http://", "https://", "data:")):
        return None
    candidate = (source_path.parent / src).resolve()
    if candidate.is_file():
        return candidate
    alt_cand = Path(src).resolve()
    if alt_cand.is_file():
        return alt_cand
    if candidate.suffix.lower() == ".svg" and candidate.with_suffix(".png").is_file():
        return candidate.with_suffix(".png")
    return None


def _embed_image_picture(paragraph: "Paragraph", image_path: Path, options: DocxBuildOptions) -> bool:
    """Embed local picture file into paragraph run."""
    try:
        from docx.image.image import Image as DocxImage

        img = DocxImage.from_file(str(image_path))
        target_w = _calculate_image_width(img.px_width, options)
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(target_w))
        return True
    except Exception:
        return False


def _add_image_to_paragraph(
    paragraph: "Paragraph",
    src: str,
    alt: str,
    options: DocxBuildOptions,
) -> bool:
    """Resolve and embed image into paragraph or add fallback text."""
    resolved = _resolve_image_file(src, options.source_path)
    if resolved and _embed_image_picture(paragraph, resolved, options):
        return True
    _add_styled_run(
        paragraph,
        f"[تصویر: {alt or src}]",
        font_name=options.font_family,
        font_size_pt=options.font_size_pt,
        italic=True,
        color=RGBColor(0x6B, 0x72, 0x80),
        rtl=options.rtl,
    )
    return False


def _add_standalone_image(
    document: "DocxDocument",
    src: str,
    alt: str,
    options: DocxBuildOptions,
) -> "Paragraph":
    """Add a standalone centered paragraph containing the image."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_image_to_paragraph(paragraph, src, alt, options)
    return paragraph


def _handle_html_inline_image(
    paragraph: "Paragraph",
    raw_html: str,
    options: DocxBuildOptions,
) -> None:
    match = HTML_IMG_RE.search(raw_html)
    if not match:
        return
    attrs = match.group("attrs")
    src_m = ATTR_SRC_RE.search(attrs)
    alt_m = ATTR_ALT_RE.search(attrs)
    src = src_m.group(1) if src_m else ""
    alt = alt_m.group(1) if alt_m else ""
    if src:
        _add_image_to_paragraph(paragraph, src, alt, options)


class HtmlBlockImageParser(HTMLParser):
    """Parse HTML blocks extracting image tags and caption containers."""

    def __init__(self) -> None:
        super().__init__()
        self.images: list[tuple[str, str]] = []
        self.caption_runs: list[tuple[str, str, str | None]] = []
        self._in_caption = False
        self._current_href: str | None = None
        self._ignore = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attr_dict = {k.lower(): v or "" for k, v in attrs}
        t = tag.lower()
        if t in ("style", "script"):
            self._ignore = True
        elif t == "img":
            self.images.append((attr_dict.get("src", ""), attr_dict.get("alt", "")))
        elif t in ("figcaption",) or "caption" in attr_dict.get("class", "").lower():
            self._in_caption = True
        elif t == "a":
            self._current_href = attr_dict.get("href")
        elif t == "br" and self._in_caption:
            self.caption_runs.append(("break", "", None))

    def handle_endtag(self, tag: str) -> None:
        t = tag.lower()
        if t in ("style", "script"):
            self._ignore = False
        elif t in ("figcaption", "div"):
            self._in_caption = False
        elif t == "a":
            self._current_href = None

    def handle_data(self, data: str) -> None:
        if self._ignore or not self._in_caption or not data:
            return
        if self._current_href:
            self.caption_runs.append(("link", data, self._current_href))
        else:
            self.caption_runs.append(("text", data, None))


def _append_caption_runs(
    paragraph: "Paragraph",
    caption_runs: list[tuple[str, str, str | None]],
    options: DocxBuildOptions,
) -> None:
    c_size = max(8, options.font_size_pt - 2)
    c_color = RGBColor(0x4B, 0x55, 0x63)
    at_line_start = True
    for kind, text, href in caption_runs:
        clean = re.sub(r"[ \t\r\n]+", " ", text)
        if kind == "break":
            paragraph.add_run().add_break()
            at_line_start = True
        elif kind == "link" and href and clean.strip():
            _add_hyperlink_run(paragraph, clean.strip(), href, font_name=options.font_family, font_size_pt=c_size, rtl=options.rtl)
            at_line_start = False
        elif kind == "text" and clean.strip():
            val = clean.lstrip() if at_line_start else clean
            _add_styled_run(paragraph, val, font_name=options.font_family, font_size_pt=c_size, color=c_color, rtl=options.rtl)
            at_line_start = False


def _render_caption_paragraph(
    document: "DocxDocument",
    caption_runs: list[tuple[str, str, str | None]],
    options: DocxBuildOptions,
) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if options.rtl:
        _set_paragraph_rtl(p)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    _append_caption_runs(p, caption_runs, options)


def _render_html_images_block(
    document: "DocxDocument",
    raw_html: str,
    options: DocxBuildOptions,
) -> None:
    parser = HtmlBlockImageParser()
    parser.feed(raw_html)
    for src, alt in parser.images:
        _add_standalone_image(document, src, alt, options)
    if parser.caption_runs:
        _render_caption_paragraph(document, parser.caption_runs, options)


def _handle_html_block(
    document: "DocxDocument",
    raw_html: str,
    options: DocxBuildOptions,
) -> None:
    raw_lower = raw_html.lower()
    if "pagebreak" in raw_lower or "page-break" in raw_lower:
        document.add_page_break()
    if "<img" in raw_lower:
        _render_html_images_block(document, raw_html, options)


def _render_inline_tokens(
    paragraph: "Paragraph",
    tokens: list[Token],
    options: DocxBuildOptions,
    *,
    base_size: int | None = None,
    base_bold: bool = False,
) -> None:
    """Render inline markdown tokens into runs of the paragraph."""
    bold_stack = [base_bold]
    italic_stack = [False]
    ltr_stack = [False]

    for token in tokens:
        kind = token.type

        if kind == "html_inline":
            content_lower = token.content.lower()
            if 'dir="ltr"' in content_lower or "dir='ltr'" in content_lower:
                ltr_stack.append(True)
            elif 'dir="rtl"' in content_lower or "dir='rtl'" in content_lower:
                ltr_stack.append(False)
            elif "</span" in content_lower and len(ltr_stack) > 1:
                ltr_stack.pop()
            elif "<img" in content_lower:
                _handle_html_inline_image(paragraph, token.content, options)
            continue

        if kind == "text":
            text = token.content
            if not text:
                continue
            is_rtl = options.rtl and (not ltr_stack[-1])
            link_href = getattr(paragraph, "_link_href", None)
            if link_href:
                _add_hyperlink_run(
                    paragraph,
                    text,
                    link_href,
                    font_name=options.font_family,
                    font_size_pt=base_size,
                    bold=bold_stack[-1],
                    italic=italic_stack[-1],
                    rtl=is_rtl,
                )
            else:
                _add_styled_run(
                    paragraph,
                    text,
                    font_name=options.font_family,
                    font_size_pt=base_size,
                    bold=bold_stack[-1],
                    italic=italic_stack[-1],
                    rtl=is_rtl,
                )
            paragraph._link_has_text = True  # type: ignore[attr-defined]

        elif kind == "softbreak":
            is_rtl = options.rtl and (not ltr_stack[-1])
            _add_styled_run(paragraph, " ", font_name=options.font_family, rtl=is_rtl)

        elif kind == "hardbreak":
            run = paragraph.add_run()
            run.add_break()

        elif kind == "code_inline":
            _add_styled_run(
                paragraph,
                token.content,
                font_name=options.font_family,
                font_size_pt=base_size,
                code=True,
                rtl=False,
            )
            paragraph._link_has_text = True  # type: ignore[attr-defined]

        elif kind == "strong_open":
            bold_stack.append(True)
        elif kind == "strong_close":
            if len(bold_stack) > 1:
                bold_stack.pop()

        elif kind == "em_open":
            italic_stack.append(True)
        elif kind == "em_close":
            if len(italic_stack) > 1:
                italic_stack.pop()

        elif kind == "link_open":
            href = token.attrGet("href") or ""
            # Store href in a marker
            paragraph._link_href = href  # type: ignore[attr-defined]
            paragraph._link_has_text = False  # type: ignore[attr-defined]

        elif kind == "link_close":
            paragraph._link_href = None  # type: ignore[attr-defined]
            paragraph._link_has_text = False  # type: ignore[attr-defined]

        elif kind == "image":
            alt = token.content or token.attrGet("alt") or ""
            src = token.attrGet("src") or ""
            _add_image_to_paragraph(paragraph, src, alt, options)

        elif kind == "s_open":
            # strikethrough not commonly used; ignore for now
            pass


# ─── Mermaid extraction (without HTML conversion) ──────────────────

MERMAID_FENCE_RE = re.compile(
    r"(?ms)(^|\n)(?P<fence>`{3,}|~{3,})[ \t]*mermaid[^\n]*\n(?P<code>.*?)(?:\n(?P=fence)[ \t]*(?=\n|$))"
)


def extract_mermaid_blocks(text: str) -> tuple[str, list[str]]:
    """
    Extract mermaid code blocks and replace them with placeholder fences.
    Returns (modified_text, list_of_mermaid_codes).
    """
    blocks: list[str] = []

    def repl(match: re.Match[str]) -> str:
        code = match.group("code").strip()
        idx = len(blocks)
        blocks.append(code)
        # Use a marker fence that markdown-it parses as a special code block
        return f"\n\n```__mermaid_placeholder__\n__MERMAID_PLACEHOLDER__{idx}__\n```\n\n"

    modified = MERMAID_FENCE_RE.sub(repl, text)
    return modified, blocks


def _build_renderer() -> MarkdownIt:
    renderer = MarkdownIt("default", {"html": True, "breaks": True, "typographer": True})
    renderer.enable("table")
    renderer.enable("strikethrough")
    return renderer


# ─── Block rendering ───────────────────────────────────────────────


def _add_heading(
    document: "DocxDocument",
    level: int,
    inline_token: Token,
    options: DocxBuildOptions,
) -> None:
    paragraph = document.add_paragraph()
    if options.rtl:
        _set_paragraph_rtl(paragraph)
        # Don't set jc=right; bidi handles RTL alignment in Word

    size = HEADING_SIZES.get(level, 12)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)

    _render_inline_tokens(
        paragraph,
        inline_token.children or [],
        options,
        base_size=size,
        base_bold=True,
    )


def _add_paragraph_block(
    document: "DocxDocument",
    inline_token: Token,
    options: DocxBuildOptions,
) -> None:
    children = inline_token.children or []
    visible_children = [
        t for t in children
        if (t.type != "html_inline" or "<img" in t.content.lower())
        and not (t.type == "text" and not t.content.strip())
    ]
    if not visible_children:
        return

    paragraph = document.add_paragraph()
    has_image = any(
        t.type == "image" or (t.type == "html_inline" and "<img" in t.content.lower())
        for t in children
    )
    if has_image:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif options.rtl:
        _set_paragraph_rtl(paragraph)

    _render_inline_tokens(
        paragraph,
        children,
        options,
        base_size=options.font_size_pt,
    )


CALLOUT_CONFIGS = {
    "NOTE": {"color": RGBColor(0x25, 0x63, 0xEB), "hex": "2563EB", "bg": "EFF6FF", "title": "💡 نکته:", "title_plain": "نکته:"},
    "TIP": {"color": RGBColor(0x16, 0xA3, 0x4A), "hex": "16A34A", "bg": "F0FDF4", "title": "💡 راهنما:", "title_plain": "راهنما:"},
    "IMPORTANT": {"color": RGBColor(0x7C, 0x3A, 0xED), "hex": "7C3AED", "bg": "F5F3FF", "title": "📌 مهم:", "title_plain": "مهم:"},
    "WARNING": {"color": RGBColor(0xD9, 0x77, 0x06), "hex": "D97706", "bg": "FFFBEB", "title": "⚠️ هشدار:", "title_plain": "هشدار:"},
    "CAUTION": {"color": RGBColor(0xDC, 0x26, 0x26), "hex": "DC2626", "bg": "FEF2F2", "title": "🛑 احتیاط:", "title_plain": "احتیاط:"},
}


def _add_code_block(
    document: "DocxDocument",
    code: str,
    options: DocxBuildOptions,
    lang: str = "",
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_ltr(paragraph)
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.right_indent = Inches(0.2)

    # Add shading via XML
    pPr = paragraph._p.get_or_add_pPr()
    shd = pPr.makeelement(
        qn("w:shd"),
        {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "F8FAFC"},
    )
    pPr.append(shd)

    # Attempt Pygments syntax highlighting
    tokens = None
    clean_lang = (lang or "").strip().lower().split()[0] if (lang or "").strip() else ""
    if options.highlight_code and clean_lang:
        try:
            from pygments.lexers import get_lexer_by_name

            lexer = get_lexer_by_name(clean_lang, stripall=False)
            tokens = list(lexer.get_tokens(code.rstrip("\n")))
        except Exception:
            tokens = None

    if tokens:
        from pygments.token import Comment, Keyword, Name, Number, Operator, String

        for ttype, val in tokens:
            if not val:
                continue
            color = None
            bold = False
            italic = False
            if ttype in Keyword or Keyword in ttype.split():
                color = RGBColor(0x00, 0x70, 0x20)
                bold = True
            elif ttype in String or String in ttype.split():
                color = RGBColor(0xBA, 0x21, 0x21)
            elif ttype in Comment or Comment in ttype.split():
                color = RGBColor(0x60, 0xA0, 0xB0)
                italic = True
            elif ttype in Number or Number in ttype.split():
                color = RGBColor(0x20, 0x80, 0x50)
            elif ttype in Name.Function or ttype in Name.Class:
                color = RGBColor(0x06, 0x28, 0x7E)
                bold = True
            elif ttype in Name.Builtin:
                color = RGBColor(0x00, 0x70, 0x20)
            elif ttype in Operator:
                color = RGBColor(0x55, 0x55, 0x55)

            _add_styled_run(
                paragraph,
                val,
                font_name=options.font_family,
                font_size_pt=options.font_size_pt - 1,
                bold=bold,
                italic=italic,
                color=color,
                code=True,
                rtl=False,
            )
    else:
        _add_styled_run(
            paragraph,
            code.rstrip("\n"),
            font_name=options.font_family,
            font_size_pt=options.font_size_pt - 1,
            code=True,
            rtl=False,
        )


def _add_blockquote(
    document: "DocxDocument",
    inline_tokens: list[list[Token]],
    options: DocxBuildOptions,
) -> None:
    callout_type = None
    if inline_tokens and inline_tokens[0]:
        first_token = inline_tokens[0][0]
        if first_token.type == "text":
            match = re.match(
                r"^\s*\[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]\s*",
                first_token.content,
                re.IGNORECASE,
            )
            if match:
                callout_type = match.group(1).upper()
                first_token.content = first_token.content[match.end():]

    cfg = CALLOUT_CONFIGS.get(callout_type) if callout_type else None

    for idx, inline in enumerate(inline_tokens):
        if cfg and idx == 0 and not any(t.content.strip() for t in inline if t.type == "text"):
            # If the first line was only the [!NOTE] tag and is now empty, skip creating an empty line
            continue

        paragraph = document.add_paragraph()
        if options.rtl:
            _set_paragraph_rtl(paragraph)
        paragraph.paragraph_format.left_indent = Inches(0.3)
        paragraph.paragraph_format.right_indent = Inches(0.3)

        pPr = paragraph._p.get_or_add_pPr()
        if cfg:
            shd = pPr.makeelement(
                qn("w:shd"),
                {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): cfg["bg"]},
            )
            pPr.append(shd)
            pBdr = pPr.makeelement(qn("w:pBdr"), {})
            border = pPr.makeelement(
                qn("w:right" if options.rtl else "w:left"),
                {qn("w:val"): "single", qn("w:sz"): "24", qn("w:space"): "8", qn("w:color"): cfg["hex"]},
            )
            pBdr.append(border)
            pPr.append(pBdr)

            if idx == 0:
                raw_title = cfg["title_plain"] if options.strip_emojis else cfg["title"]
                suffix = " " if any(t.content.strip() for t in inline if t.type == "text") else ""
                _add_styled_run(
                    paragraph,
                    raw_title + suffix,
                    font_name=options.font_family,
                    font_size_pt=options.font_size_pt,
                    bold=True,
                    color=cfg["color"],
                    rtl=options.rtl,
                )

        _render_inline_tokens(
            paragraph,
            inline,
            options,
            base_size=options.font_size_pt,
        )

        if not cfg:
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)


def _add_list_item(
    document: "DocxDocument",
    inline_token: Token,
    options: DocxBuildOptions,
    *,
    ordered: bool,
    index: int,
    level: int,
) -> None:
    paragraph = document.add_paragraph()
    if options.rtl:
        _set_paragraph_rtl(paragraph)

    indent = Inches(0.25 * (level + 1))
    if options.rtl:
        paragraph.paragraph_format.right_indent = indent
    else:
        paragraph.paragraph_format.left_indent = indent

    bullet = f"{index}. " if ordered else "• "
    _add_styled_run(
        paragraph,
        bullet,
        font_name=options.font_family,
        font_size_pt=options.font_size_pt,
        bold=True,
        rtl=options.rtl,
    )

    _render_inline_tokens(
        paragraph,
        inline_token.children or [],
        options,
        base_size=options.font_size_pt,
    )


def _add_horizontal_rule(document: "DocxDocument") -> None:
    paragraph = document.add_paragraph()
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pPr.makeelement(
        qn("w:bottom"),
        {
            qn("w:val"): "single",
            qn("w:sz"): "6",
            qn("w:space"): "1",
            qn("w:color"): "E5E7EB",
        },
    )
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_mermaid_image(
    document: "DocxDocument",
    image_path: Path,
    options: DocxBuildOptions,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    try:
        # Read PNG dimensions to calculate proportional width
        import struct

        with open(image_path, "rb") as f:
            header = f.read(24)
        img_width = struct.unpack(">I", header[16:20])[0]
        img_height = struct.unpack(">I", header[20:24])[0]

        # PNGs are rendered at options.image_scale * 96 DPI.
        # Compute natural CSS width and clamp between min and max.
        css_width_inches = img_width / (float(options.image_scale) * 96.0)
        fmt_key = (options.page_format or "A4").strip().upper()
        w_mm, h_mm = PAGE_FORMAT_DIMENSIONS.get(fmt_key, (210.0, 297.0))
        page_w_in = (h_mm if options.landscape else w_mm) / 25.4
        margin_in = parse_docx_length(options.margin).inches
        available_w = max(1.0, page_w_in - (2 * margin_in))
        max_limit = min(options.image_max_width_inches, available_w) if options.image_max_width_inches != 6.5 else available_w
        target_width = min(
            max(css_width_inches, options.image_min_width_inches),
            max_limit,
        )

        run.add_picture(str(image_path), width=Inches(target_width))
    except Exception as exc:  # noqa: BLE001
        _add_styled_run(
            paragraph,
            f"[خطا در درج نمودار Mermaid: {exc}]",
            font_name=options.font_family,
            font_size_pt=options.font_size_pt - 1,
            italic=True,
            color=RGBColor(0xDC, 0x26, 0x26),
            rtl=options.rtl,
        )


def _add_mermaid_placeholder_text(
    document: "DocxDocument",
    code: str,
    options: DocxBuildOptions,
) -> None:
    """Fallback when Mermaid image is not available: include the source code."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_styled_run(
        paragraph,
        "[نمودار Mermaid - متن منبع:]",
        font_name=options.font_family,
        font_size_pt=options.font_size_pt - 1,
        italic=True,
        color=RGBColor(0x6B, 0x72, 0x80),
        rtl=options.rtl,
    )
    _add_code_block(document, code, options)


# ─── Table rendering ───────────────────────────────────────────────


def _add_table_from_tokens(
    document: "DocxDocument",
    tokens: list[Token],
    start_index: int,
    options: DocxBuildOptions,
) -> int:
    """
    Build a docx table from markdown-it table tokens.
    Returns index of token after table_close.
    """
    rows: list[list[list[Token]]] = []  # rows -> cells -> inline children
    is_header_row = False
    current_row: list[list[Token]] = []

    i = start_index + 1  # skip table_open
    while i < len(tokens):
        token = tokens[i]
        if token.type == "table_close":
            i += 1
            break
        elif token.type == "thead_open":
            is_header_row = True
        elif token.type == "thead_close":
            is_header_row = False
        elif token.type == "tr_open":
            current_row = []
        elif token.type == "tr_close":
            rows.append(current_row)
        elif token.type in ("th_open", "td_open"):
            # next token should be inline
            if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                current_row.append(tokens[i + 1].children or [])
                i += 1  # skip inline
        i += 1

    if not rows:
        return i

    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Light Grid Accent 1"

    # Right-to-left table layout
    if options.rtl:
        tbl = table._tbl
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is not None:
            bidiVisual = tblPr.find(qn("w:bidiVisual"))
            if bidiVisual is None:
                bidiVisual = tblPr.makeelement(qn("w:bidiVisual"), {})
                tblPr.append(bidiVisual)

    for row_idx, row_cells in enumerate(rows):
        for col_idx, inline_tokens in enumerate(row_cells):
            cell = table.rows[row_idx].cells[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            # Clear default empty paragraph
            paragraph = cell.paragraphs[0]
            paragraph.text = ""
            if options.rtl:
                _set_paragraph_rtl(paragraph)
                # Don't set jc; bidi handles alignment

            is_header = row_idx == 0
            _render_inline_tokens(
                paragraph,
                inline_tokens,
                options,
                base_size=options.font_size_pt - 1,
                base_bold=is_header,
            )

    return i


# ─── Main build function ───────────────────────────────────────────


def _is_mermaid_placeholder_block(token: Token) -> tuple[bool, int | None]:
    """Check if a fence token is a mermaid placeholder. Returns (is_placeholder, index)."""
    if token.type != "fence":
        return False, None
    if (token.info or "").strip() != "__mermaid_placeholder__":
        return False, None
    match = re.search(r"__MERMAID_PLACEHOLDER__(\d+)__", token.content)
    if not match:
        return False, None
    return True, int(match.group(1))


def build_docx(
    markdown_text: str,
    options: DocxBuildOptions,
    output_path: Path,
) -> DocxBuildResult:
    """
    Build a DOCX file from Markdown text.

    If options.mermaid_images is provided, mermaid blocks are replaced with
    embedded PNG images. Otherwise the source code is included as a code block.
    """
    if options.strip_emojis:
        markdown_text = strip_emojis(markdown_text)
    text = strip_front_matter(markdown_text)
    title = extract_title(text, options.source_path.stem)

    # Extract mermaid blocks before parsing
    text_with_placeholders, mermaid_blocks = extract_mermaid_blocks(text)

    renderer = _build_renderer()
    tokens = renderer.parse(text_with_placeholders)

    document = Document()
    _set_document_default_rtl(document, options.font_family, options.font_size_pt)

    # Set page dimensions, orientation, margins, and footer page numbers
    fmt_key = (options.page_format or "A4").strip().upper()
    w_mm, h_mm = PAGE_FORMAT_DIMENSIONS.get(fmt_key, (210.0, 297.0))
    margin_len = parse_docx_length(options.margin)

    for section in document.sections:
        if options.landscape:
            section.page_width = Mm(h_mm)
            section.page_height = Mm(w_mm)
            section.orientation = WD_ORIENT.LANDSCAPE
        else:
            section.page_width = Mm(w_mm)
            section.page_height = Mm(h_mm)
            section.orientation = WD_ORIENT.PORTRAIT

        section.left_margin = margin_len
        section.right_margin = margin_len
        section.top_margin = margin_len
        section.bottom_margin = margin_len
        # RTL section - set bidi on section properties
        if options.rtl:
            sectPr = section._sectPr
            bidi = sectPr.find(qn("w:bidi"))
            if bidi is None:
                bidi = sectPr.makeelement(qn("w:bidi"), {})
                sectPr.append(bidi)
            # Also set document grid for RTL
            docGrid = sectPr.find(qn("w:docGrid"))
            if docGrid is None:
                docGrid = sectPr.makeelement(qn("w:docGrid"), {})
                sectPr.append(docGrid)
            docGrid.set(qn("w:charSpace"), "0")

        # Dynamic Persian footer page numbering: "صفحه X از Y"
        if options.include_page_numbers:
            footer = section.footer
            if footer.paragraphs:
                f_p = footer.paragraphs[0]
                f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if options.rtl:
                    _set_paragraph_rtl(f_p)
                r_page = f_p.add_run("صفحه ")
                _set_run_font(r_page, options.font_family)
                r_page.font.size = Pt(9)
                r_page.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

                from docx.oxml import parse_xml

                f_p.add_run()._r.append(
                    parse_xml(r'<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="PAGE"/>')
                )
                r_of = f_p.add_run(" از ")
                _set_run_font(r_of, options.font_family)
                r_of.font.size = Pt(9)
                r_of.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

                f_p.add_run()._r.append(
                    parse_xml(r'<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="NUMPAGES"/>')
                )

    # Walk tokens
    i = 0
    list_levels: list[bool] = []  # stack of ordered flags
    list_indices: list[int] = []  # stack of current index per list

    while i < len(tokens):
        token = tokens[i]
        kind = token.type

        if kind.startswith("heading_open"):
            level = int(kind[-1]) if kind[-1].isdigit() else int(token.tag[1:])
            if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                _add_heading(document, level, tokens[i + 1], options)
            i += 3  # heading_open, inline, heading_close
            continue

        if kind == "paragraph_open":
            if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                inline = tokens[i + 1]
                # Check if inside a list item
                if list_levels:
                    ordered = list_levels[-1]
                    list_indices[-1] += 1
                    _add_list_item(
                        document,
                        inline,
                        options,
                        ordered=ordered,
                        index=list_indices[-1],
                        level=len(list_levels) - 1,
                    )
                else:
                    _add_paragraph_block(document, inline, options)
            i += 3  # paragraph_open, inline, paragraph_close
            continue

        if kind == "fence" or kind == "code_block":
            is_mermaid, mermaid_idx = _is_mermaid_placeholder_block(token)
            if is_mermaid and mermaid_idx is not None and mermaid_idx < len(mermaid_blocks):
                mermaid_code = mermaid_blocks[mermaid_idx]
                if options.mermaid_images and mermaid_code in options.mermaid_images:
                    _add_mermaid_image(
                        document, options.mermaid_images[mermaid_code], options
                    )
                else:
                    _add_mermaid_placeholder_text(document, mermaid_code, options)
            else:
                lang = (token.info or "").strip().split()[0] if (token.info or "").strip() else ""
                _add_code_block(document, token.content, options, lang=lang)
            i += 1
            continue

        if kind == "bullet_list_open":
            list_levels.append(False)
            list_indices.append(0)
            i += 1
            continue
        if kind == "ordered_list_open":
            list_levels.append(True)
            list_indices.append(0)
            i += 1
            continue
        if kind in ("bullet_list_close", "ordered_list_close"):
            if list_levels:
                list_levels.pop()
                list_indices.pop()
            i += 1
            continue

        if kind in ("list_item_open", "list_item_close"):
            i += 1
            continue

        if kind == "blockquote_open":
            inlines: list[list[Token]] = []
            j = i + 1
            depth = 1
            while j < len(tokens) and depth > 0:
                if tokens[j].type == "blockquote_open":
                    depth += 1
                elif tokens[j].type == "blockquote_close":
                    depth -= 1
                if depth == 0:
                    break
                elif tokens[j].type == "inline":
                    inlines.append(tokens[j].children or [])
                j += 1
            _add_blockquote(document, inlines, options)
            i = j + 1
            continue

        if kind == "hr":
            _add_horizontal_rule(document)
            i += 1
            continue

        if kind == "table_open":
            i = _add_table_from_tokens(document, tokens, i, options)
            continue

        if kind == "html_block":
            _handle_html_block(document, token.content, options)
            i += 1
            continue

        # Skip unknown tokens
        i += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))

    return DocxBuildResult(
        title=title,
        has_mermaid=bool(mermaid_blocks),
        mermaid_blocks=mermaid_blocks,
    )
