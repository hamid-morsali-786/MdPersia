from pathlib import Path
import docx
from fa_md_pdf.docx_builder import DocxBuildOptions, build_docx


def test_docx_omits_html_and_sets_code_ltr(tmp_path: Path) -> None:
    md_content = """<div dir="rtl">
<style>body { font-size: 14px; }</style>

# عنوان تستی

متن آزمایشی دارای <span dir="ltr">`(Oracle RAC)`</span> است.

```sql
SELECT 1 FROM DUAL;
```
</div>
"""
    input_file = tmp_path / "test.md"
    input_file.write_text(md_content, encoding="utf-8")
    output_file = tmp_path / "output.docx"

    options = DocxBuildOptions(source_path=input_file)
    build_docx(md_content, options, output_file)

    assert output_file.is_file()
    doc = docx.Document(output_file)

    # 1. No raw HTML tags should be present
    for p in doc.paragraphs:
        assert "<div" not in p.text
        assert "<style" not in p.text
        assert "<span" not in p.text

    # 2. Code block should be LTR
    code_p = [p for p in doc.paragraphs if "SELECT 1" in p.text]
    assert len(code_p) == 1
    bidi_nodes = code_p[0]._p.xpath('.//w:bidi')
    assert len(bidi_nodes) > 0
    assert bidi_nodes[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "0"


def test_docx_embeds_local_image(tmp_path: Path) -> None:
    import base64

    # 1x1 valid PNG
    b64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
    png_bytes = base64.b64decode(b64)
    img_file = tmp_path / "sample.png"
    img_file.write_bytes(png_bytes)

    md_content = f"![نمونه]({img_file.name})"
    input_file = tmp_path / "doc.md"
    input_file.write_text(md_content, encoding="utf-8")
    output_file = tmp_path / "doc.docx"

    options = DocxBuildOptions(source_path=input_file)
    build_docx(md_content, options, output_file)

    assert output_file.is_file()
    doc = docx.Document(output_file)

    # Verify drawing element was embedded
    drawings = []
    for p in doc.paragraphs:
        drawings.extend(p._p.xpath('.//w:drawing'))
    assert len(drawings) == 1


def test_docx_custom_margins_and_page_format(tmp_path: Path) -> None:
    from docx.enum.section import WD_ORIENT

    md_content = "# تست حاشیه\nمتن نمونه"
    input_file = tmp_path / "margin_test.md"
    input_file.write_text(md_content, encoding="utf-8")
    output_file = tmp_path / "margin_test.docx"

    options = DocxBuildOptions(
        source_path=input_file,
        margin="5mm",
        page_format="A4",
        landscape=True,
    )
    build_docx(md_content, options, output_file)

    assert output_file.is_file()
    doc = docx.Document(output_file)
    section = doc.sections[0]

    # Verify 5mm margins (approx 4.99mm due to twips conversion)
    assert abs(section.left_margin.mm - 5.0) < 0.1
    assert abs(section.right_margin.mm - 5.0) < 0.1
    assert abs(section.top_margin.mm - 5.0) < 0.1
    assert abs(section.bottom_margin.mm - 5.0) < 0.1

    # Verify landscape A4 (297mm width, 210mm height)
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert abs(section.page_width.mm - 297.0) < 0.5
    assert abs(section.page_height.mm - 210.0) < 0.5
