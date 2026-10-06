"""
End-to-End (E2E) integration pipeline tests for fa-md-pdf.
Executes the full pipeline converting realistic Persian documents with Mermaid diagrams,
callouts, tables, and code highlighting to both PDF (via Playwright) and DOCX (via python-docx).
"""

from __future__ import annotations

from pathlib import Path
import docx
import pytest

from fa_md_pdf.converter import ConvertJob, ConvertOptions, convert_jobs
from fa_md_pdf.rtl_wrapper import wrap_rtl_in_file


def test_e2e_full_pipeline_pdf(
    tmp_path: Path,
    complex_persian_markdown: str,
    vendor_mermaid: Path,
):
    if not vendor_mermaid.is_file():
        pytest.skip("Offline mermaid.min.js not found in vendor")

    doc_file = tmp_path / "e2e_test.md"
    doc_file.write_text(complex_persian_markdown, encoding="utf-8")
    pdf_out = tmp_path / "e2e_test.pdf"
    html_out = tmp_path / "e2e_test.html"

    job = ConvertJob(source=doc_file, output=pdf_out, html_output=html_out)
    options = ConvertOptions(
        mermaid_source=str(vendor_mermaid),
        keep_html=True,
        verbose=True,
        mermaid_timeout_ms=30_000,
    )

    results = convert_jobs([job], options=options)
    assert len(results) == 1
    assert results[0].ok is True
    assert results[0].error is None

    # Verify generated PDF
    assert pdf_out.is_file()
    pdf_bytes = pdf_out.read_bytes()
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 10_000

    # Verify intermediate HTML
    assert html_out.is_file()
    html_text = html_out.read_text(encoding="utf-8")
    assert '<main class="markdown-body">' in html_text
    assert '<pre class="mermaid">' in html_text
    assert "class=\"callout callout-note\"" in html_text


def test_e2e_full_pipeline_docx(
    tmp_path: Path,
    complex_persian_markdown: str,
    vendor_mermaid: Path,
):
    if not vendor_mermaid.is_file():
        pytest.skip("Offline mermaid.min.js not found in vendor")

    doc_file = tmp_path / "e2e_test.md"
    doc_file.write_text(complex_persian_markdown, encoding="utf-8")
    docx_out = tmp_path / "e2e_test.docx"

    job = ConvertJob(source=doc_file, output=docx_out)
    options = ConvertOptions(
        output_format="docx",
        mermaid_source=str(vendor_mermaid),
        verbose=True,
        mermaid_timeout_ms=30_000,
    )

    results = convert_jobs([job], options=options)
    assert len(results) == 1
    assert results[0].ok is True
    assert results[0].error is None

    # Verify generated DOCX
    assert docx_out.is_file()
    assert docx_out.stat().st_size > 10_000

    doc = docx.Document(docx_out)
    # Check paragraphs, callout boxes, and table
    assert len(doc.paragraphs) > 10
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 4

    # Verify that the Mermaid diagram was rendered and embedded as a drawing
    drawings = []
    for p in doc.paragraphs:
        drawings.extend(p._p.xpath(".//w:drawing"))
    assert len(drawings) >= 1


def test_e2e_wrap_rtl_then_convert(
    tmp_path: Path,
    complex_persian_markdown: str,
    vendor_mermaid: Path,
):
    if not vendor_mermaid.is_file():
        pytest.skip("Offline mermaid.min.js not found in vendor")

    raw_file = tmp_path / "raw.md"
    raw_file.write_text(complex_persian_markdown, encoding="utf-8")

    # Step 1: Wrap with RTL divs
    wrapped_file = wrap_rtl_in_file(raw_file)
    assert wrapped_file.is_file()
    assert '<div dir="rtl">' in wrapped_file.read_text(encoding="utf-8")

    # Step 2: Convert wrapped document to DOCX
    docx_out = tmp_path / "wrapped.docx"
    job = ConvertJob(source=wrapped_file, output=docx_out)
    options = ConvertOptions(
        output_format="docx",
        mermaid_source=str(vendor_mermaid),
    )
    results = convert_jobs([job], options=options)
    assert len(results) == 1
    assert results[0].ok is True
    assert docx_out.is_file()
