"""
Advanced tests for DOCX generation: headings, lists, tables, callouts,
code highlighting, fallback handling, mixed bidi, and page setups.
"""

from __future__ import annotations

from pathlib import Path
import docx
import pytest

from fa_md_pdf.docx_builder import (
    DocxBuildOptions,
    build_docx,
    extract_mermaid_blocks,
    parse_docx_length,
)


def test_docx_headings_hierarchy(tmp_path: Path):
    md = """# سرتیتر سطح ۱
## سرتیتر سطح ۲
### سرتیتر سطح ۳
#### سرتیتر سطح ۴
##### سرتیتر سطح ۵
###### سرتیتر سطح ۶
"""
    output = tmp_path / "headings.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "h.md"), output)

    doc = docx.Document(output)
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    assert len(paragraphs) == 6

    for p in paragraphs:
        # Every heading in RTL mode must have bidi element
        assert len(p._p.xpath(".//w:bidi")) > 0
        # Check run exists
        assert len(p.runs) >= 1
        assert p.runs[0].bold is True


def test_docx_unordered_and_ordered_lists(tmp_path: Path):
    md = """# لیست‌ها

- مورد نشانه‌دار ۱
- مورد نشانه‌دار ۲
  - زیرمورد تو رفته

1. مورد شماره‌دار ۱
2. مورد شماره‌دار ۲
"""
    output = tmp_path / "lists.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "lists.md"), output)

    doc = docx.Document(output)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]

    assert any("مورد نشانه‌دار ۱" in t for t in texts)
    assert any("مورد شماره‌دار ۱" in t for t in texts)
    assert any("زیرمورد تو رفته" in t for t in texts)


def test_docx_table_styling_and_rtl_cells(tmp_path: Path):
    md = """| ستون اول | ستون دوم | ستون سوم |
| :--- | :---: | ---: |
| داده ۱ | داده ۲ | داده ۳ |
| مقدار الف | مقدار ب | مقدار ج |
"""
    output = tmp_path / "table.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "tbl.md"), output)

    doc = docx.Document(output)
    assert len(doc.tables) == 1

    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 3

    # Verify table properties have RTL direction
    tblPr = table._tbl.tblPr
    assert len(tblPr.xpath(".//w:bidiVisual")) > 0

    # Table style should be Light Grid Accent 1
    assert table.style.name == "Light Grid Accent 1"

    # Header cells must contain bold text
    for cell in table.rows[0].cells:
        assert any(r.bold is True for r in cell.paragraphs[0].runs)


def test_docx_all_callout_types(tmp_path: Path):
    md = """
> [!NOTE]
> یادداشت تستی

> [!TIP]
> راهنمای تستی

> [!IMPORTANT]
> نکته مهم تستی

> [!WARNING]
> هشدار تستی

> [!CAUTION]
> احتیاط تستی

> نقل‌قول معمولی بدون تگ کال‌اوت.
"""
    output = tmp_path / "callouts.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "c.md"), output)

    doc = docx.Document(output)
    full_text = " ".join(p.text for p in doc.paragraphs)

    assert "💡 نکته:" in full_text
    assert "💡 راهنما:" in full_text
    assert "📌 مهم:" in full_text
    assert "⚠️ هشدار:" in full_text
    assert "🛑 احتیاط:" in full_text
    assert "نقل‌قول معمولی" in full_text


def test_docx_code_block_highlighted_and_unhighlighted(tmp_path: Path):
    md = """```python
def add(a, b):
    # جمع دو عدد
    return a + b
```

```text
متن خام بدون زبان برنامه‌نویسی
```
"""
    output = tmp_path / "code.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "code.md"), output)

    doc = docx.Document(output)
    # Check that code block paragraphs are LTR
    code_paragraphs = [p for p in doc.paragraphs if "def add" in p.text or "متن خام" in p.text]
    assert len(code_paragraphs) >= 2

    for cp in code_paragraphs:
        bidi = cp._p.xpath(".//w:bidi")
        if bidi:
            assert bidi[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "0"


def test_docx_mermaid_embedding_and_fallback(tmp_path: Path, sample_png_bytes: bytes):
    png_path = tmp_path / "rendered_diagram.png"
    png_path.write_bytes(sample_png_bytes)

    diagram_with_img = "graph TD\n    A --> B"
    diagram_without_img = "graph LR\n    X --> Y"

    md = f"""# دیاگرام‌ها

```mermaid
{diagram_with_img}
```

```mermaid
{diagram_without_img}
```
"""
    output = tmp_path / "mermaid.docx"
    options = DocxBuildOptions(
        source_path=tmp_path / "doc.md",
        mermaid_images={diagram_with_img: png_path},  # diagram_without_img is omitted
    )
    build_docx(md, options, output)

    doc = docx.Document(output)
    # 1. Embedded image drawing check
    drawings = []
    for p in doc.paragraphs:
        drawings.extend(p._p.xpath(".//w:drawing"))
    assert len(drawings) == 1

    # 2. Fallback code block check for diagram_without_img
    fallback_p = [p for p in doc.paragraphs if "X --> Y" in p.text]
    assert len(fallback_p) == 1


def test_docx_missing_image_handled_gracefully(tmp_path: Path):
    md = """# تست تصویر ناموجود
![تصویر گمشده](non_existent_image_12345.png)
متن پس از تصویر.
"""
    output = tmp_path / "missing_img.docx"
    # Must not raise an unhandled exception when an image file is missing
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "doc.md"), output)
    assert output.is_file()
    doc = docx.Document(output)
    assert any("متن پس از تصویر" in p.text for p in doc.paragraphs)


def test_docx_empty_and_whitespace_markdown(tmp_path: Path):
    output_empty = tmp_path / "empty.docx"
    build_docx("", DocxBuildOptions(source_path=tmp_path / "empty.md"), output_empty)
    assert output_empty.is_file()

    output_ws = tmp_path / "ws.docx"
    build_docx("   \n\n\t  \n", DocxBuildOptions(source_path=tmp_path / "ws.md"), output_ws)
    assert output_ws.is_file()


def test_docx_mixed_bidi_and_inline_formatting(tmp_path: Path):
    md = """# متن ترکیبی

متن دارای **بولد فارسی** و *ایتالیک* و ~~خط‌خورده~~ و `کد درون‌خطی InlineCode` است.
عبارت انگلیسی (Enterprise Architecture Framework) در متن فارسی.
"""
    output = tmp_path / "mixed.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "mixed.md"), output)

    doc = docx.Document(output)
    p = [p for p in doc.paragraphs if "متن دارای" in p.text][0]

    # Verify bold, italic, and code runs exist
    has_bold = any(r.bold for r in p.runs)
    has_italic = any(r.italic for r in p.runs)
    has_code = any("InlineCode" in r.text for r in p.runs)

    assert has_bold is True
    assert has_italic is True
    assert has_code is True


def test_docx_page_dimensions_letter_and_legal(tmp_path: Path):
    from docx.enum.section import WD_ORIENT

    md = "# سند"
    out_letter = tmp_path / "letter.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "l.md", page_format="Letter"), out_letter)
    doc_letter = docx.Document(out_letter)
    assert abs(doc_letter.sections[0].page_width.mm - 215.9) < 0.5

    out_legal = tmp_path / "legal.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "lg.md", page_format="Legal"), out_legal)
    doc_legal = docx.Document(out_legal)
    assert abs(doc_legal.sections[0].page_height.mm - 355.6) < 0.5


def test_docx_callout_with_blank_line_and_standalone(tmp_path: Path):
    # 1. Callout where tag is followed by blank line
    md_multiline = "> [!NOTE]\n>\n> متن یادداشت در پاراگراف دوم\n"
    out_multi = tmp_path / "callout_multi.docx"
    build_docx(md_multiline, DocxBuildOptions(source_path=tmp_path / "c1.md"), out_multi)
    doc_multi = docx.Document(out_multi)
    full_text_multi = " ".join(p.text for p in doc_multi.paragraphs)
    assert "💡 نکته:" in full_text_multi
    assert "متن یادداشت در پاراگراف دوم" in full_text_multi

    # 2. Standalone callout with no subsequent text
    md_standalone = "> [!TIP]\n"
    out_solo = tmp_path / "callout_solo.docx"
    build_docx(md_standalone, DocxBuildOptions(source_path=tmp_path / "c2.md"), out_solo)
    doc_solo = docx.Document(out_solo)
    assert len(doc_solo.paragraphs) >= 1
    assert any("💡 راهنما:" in p.text for p in doc_solo.paragraphs)


def test_docx_cleans_surrogates_and_xml_noncharacters(tmp_path: Path):
    # Text containing XML 1.0 illegal noncharacters \ufffe, \uffff and control chars
    md = "# سرتیتر آزمایشی \ufffe\n\nمتن حاوی کاراکترهای \uffff غیرمجاز و کنترل \x00\x0c است."
    out_docx = tmp_path / "nonchars.docx"
    build_docx(md, DocxBuildOptions(source_path=tmp_path / "nonchars.md"), out_docx)
    assert out_docx.is_file()

    doc = docx.Document(out_docx)
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "\ufffe" not in full_text
    assert "\uffff" not in full_text
    assert "\x00" not in full_text
    assert "سرتیتر آزمایشی" in full_text
    assert "غیرمجاز" in full_text


def test_docx_empty_table_guard(tmp_path: Path):
    from fa_md_pdf.docx_builder import _add_table_from_tokens
    from markdown_it.token import Token

    # Construct degenerate table tokens where rows have 0 cells
    tokens = [
        Token(type="table_open", tag="table", nesting=1),
        Token(type="tr_open", tag="tr", nesting=1),
        Token(type="tr_close", tag="tr", nesting=-1),
        Token(type="table_close", tag="table", nesting=-1),
    ]
    doc = docx.Document()
    idx = _add_table_from_tokens(doc, tokens, 0, DocxBuildOptions(source_path=tmp_path / "doc.md"))
    assert idx == len(tokens)
    # No degenerate 0-column table should be added to the document
    assert len(doc.tables) == 0
