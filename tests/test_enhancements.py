from pathlib import Path
import docx
from fa_md_pdf.docx_builder import DocxBuildOptions, build_docx
from fa_md_pdf.html_builder import HtmlBuildOptions, build_html


def test_html_syntax_highlighting_and_callouts(tmp_path: Path) -> None:
    md_content = """# عنوان

> [!NOTE]
> این یک یادداشت مهم است.

```python
def hello_world():
    return "سلام"
```

<!-- pagebreak -->

پایان صفحه اول.
"""
    doc_path = tmp_path / "test.md"
    doc_path.write_text(md_content, encoding="utf-8")

    html_doc = build_html(
        md_content,
        HtmlBuildOptions(source_path=doc_path),
    )

    # 1. Check callout transformed
    assert 'class="callout callout-note"' in html_doc.html
    assert "💡 نکته" in html_doc.html

    # 2. Check syntax highlighting
    # 'def' is a keyword, 'return' is a keyword
    assert 'class="k">def</span>' in html_doc.html or 'class="k">return</span>' in html_doc.html

    # 3. Check pagebreak transformed
    assert '<div class="page-break"></div>' in html_doc.html


def test_docx_hyperlink_callout_and_footer(tmp_path: Path) -> None:
    md_content = """# تست ورد

> [!WARNING]
> اخطار به سیستم!

اینجا یک [لینک رسمی](https://example.com) وجود دارد.

```sql
SELECT ID, NAME FROM USERS WHERE ACTIVE = 1;
```

<!-- pagebreak -->

صفحه دوم سند.
"""
    doc_path = tmp_path / "test_docx.md"
    doc_path.write_text(md_content, encoding="utf-8")
    out_docx = tmp_path / "test_docx.docx"

    build_docx(
        md_content,
        DocxBuildOptions(source_path=doc_path),
        out_docx,
    )

    assert out_docx.is_file()
    doc = docx.Document(out_docx)

    # 1. Native hyperlink check
    hyperlinks = doc.paragraphs[1]._p.xpath('.//w:hyperlink') or doc.paragraphs[2]._p.xpath('.//w:hyperlink')
    assert len(hyperlinks) >= 1

    # 2. Callout box in Word check
    callout_p = [p for p in doc.paragraphs if "⚠️ هشدار:" in p.text]
    assert len(callout_p) == 1

    # 3. Dynamic footer check
    footer = doc.sections[0].footer
    footer_fields = footer.paragraphs[0]._p.xpath('.//w:fldSimple')
    instr_list = [f.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr") for f in footer_fields]
    assert "PAGE" in instr_list
    assert "NUMPAGES" in instr_list
