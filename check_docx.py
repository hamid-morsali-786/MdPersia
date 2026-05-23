"""Check RTL settings in generated DOCX."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

doc = Document(str(Path("output/GUI-PROPOSAL-v2.docx")))

# Check Normal style
style = doc.styles["Normal"]
pPr = style.element.find(qn("w:pPr"))
print("=== Normal Style ===")
print(f"  pPr exists: {pPr is not None}")
if pPr is not None:
    bidi = pPr.find(qn("w:bidi"))
    jc = pPr.find(qn("w:jc"))
    print(f"  bidi: {bidi is not None}")
    print(f"  jc val: {jc.get(qn('w:val')) if jc is not None else 'MISSING'}")

rPr = style.element.find(qn("w:rPr"))
if rPr is not None:
    rtl = rPr.find(qn("w:rtl"))
    print(f"  rPr rtl: {rtl is not None}")

# Check section
print("\n=== Section ===")
sect = doc.sections[0]._sectPr
bidi = sect.find(qn("w:bidi"))
print(f"  bidi: {bidi is not None}")
print(f"  Section XML snippet:")
print(f"  {etree.tostring(sect, pretty_print=True).decode()[:500]}")

# Check paragraph 1 (title) in detail
print("\n=== Paragraph 1 (Title) ===")
p = doc.paragraphs[1]
print(f"  Text: {p.text[:60]}")
ppPr = p._p.find(qn("w:pPr"))
if ppPr is not None:
    print(f"  pPr XML: {etree.tostring(ppPr).decode()}")
if p.runs:
    r = p.runs[0]
    rrPr = r._r.find(qn("w:rPr"))
    if rrPr is not None:
        print(f"  Run rPr XML: {etree.tostring(rrPr).decode()}")

# Check paragraph 3 (body text)
print("\n=== Paragraph 3 (Body) ===")
p = doc.paragraphs[3]
print(f"  Text: {p.text[:60]}")
ppPr = p._p.find(qn("w:pPr"))
if ppPr is not None:
    print(f"  pPr XML: {etree.tostring(ppPr).decode()}")
